"""Document redaction: every format, every hiding place, hostile files and the process boundary."""

from __future__ import annotations

import asyncio
import io
import json
import os
import re
import signal
import zipfile
from types import SimpleNamespace
from typing import Any

import pytest

from harpo_redact.documents.runner import DocumentError, WorkerRunner
from harpo_redact.documents.service import DocumentOutcome, redact_document

from . import doc_fixtures as fx
from .conftest import make_settings


@pytest.fixture(scope="module")
def env(ner):  # type: ignore[no-untyped-def]
    from harpo_redact.pipeline import Pipeline

    def build(**overrides: Any) -> SimpleNamespace:
        settings = make_settings(**overrides)
        return SimpleNamespace(
            settings=settings, pipeline=Pipeline(settings, ner=ner), runner=WorkerRunner(settings.documents)
        )

    return build


def run(e: SimpleNamespace, data: bytes, name: str, content_type: str | None = None, **kw: Any) -> DocumentOutcome:
    return asyncio.run(
        redact_document(
            e.pipeline, e.settings.documents, e.runner, data, name, content_type,
            mode=kw.get("mode"), context_repair=None, override=kw.get("override", False), want_file=kw.get("want_file", True),
        )
    )  # fmt: skip


def refused(e: SimpleNamespace, data: bytes, name: str, **kw: Any) -> str:
    with pytest.raises(DocumentError) as info:
        run(e, data, name, **kw)
    return info.value.code


def with_part(data: bytes, name: str, blob: bytes | None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename != name:
                dst.writestr(info.filename, src.read(info.filename))
        if blob is not None:
            dst.writestr(name, blob)
    return out.getvalue()


@pytest.fixture(scope="module")
def docx_run(env):  # type: ignore[no-untyped-def]
    e = env()
    return e, fx.make_docx(), run(e, fx.make_docx(), "Priya Raman CV.docx")


@pytest.fixture(scope="module")
def xlsx_run(env):  # type: ignore[no-untyped-def]
    return run(env(), fx.make_xlsx(), "customers.xlsx")


@pytest.fixture(scope="module")
def pptx_run(env):  # type: ignore[no-untyped-def]
    return run(env(), fx.make_pptx(), "deck.pptx")


@pytest.fixture(scope="module")
def pdf_run(env):  # type: ignore[no-untyped-def]
    e = env()
    data = fx.make_pdf()
    return data, run(e, data, "invoice.pdf")


# =========================================================================== Word
class TestDocx:
    def test_nothing_survives_in_any_part(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, _, out = docx_run
        assert out.verified and out.data is not None
        assert fx.canaries_in_package(out.data) == []

    def test_nothing_survives_for_an_independent_reader(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, _, out = docx_run
        assert out.data is not None
        view = fx.docx_text(out.data)
        joined = "\n".join(view["texts"] + view["meta"])
        for value in fx.CANARIES:
            assert value not in joined, value

    def test_split_run_name_becomes_one_placeholder_and_formatting_survives(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, original, out = docx_run
        assert out.data is not None
        before, after = fx.docx_text(original)["doc"], fx.docx_text(out.data)["doc"]
        assert after.paragraphs[0].text == "Please contact [PERSON_1] at [EMAIL_1] today."
        assert [len(p.runs) for p in before.paragraphs[:3]] == [len(p.runs) for p in after.paragraphs[:3]]  # runs kept
        bold = [r.text for r in after.paragraphs[2].runs if r.bold]
        assert bold == ["bold words"]  # untouched formatting
        assert len(before.tables[0].rows) == len(after.tables[0].rows)

    def test_placeholders_are_consistent_across_parts(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, _, out = docx_run
        assert out.data is not None
        parts = fx.package_text(out.data)
        # Lars Eriksen is in the table, the header and the comment: one person, one placeholder
        for part in ("word/document.xml", "word/header1.xml", "word/comments.xml"):
            assert b"[PERSON_2]" in parts[part], part
        assert b"[PERSON_1]" in parts["word/document.xml"] and b"[PERSON_1]" not in parts["word/header1.xml"]

    def test_hidden_places_were_reached(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, _, out = docx_run
        assert out.data is not None
        parts = fx.package_text(out.data)
        body = parts["word/document.xml"].decode()
        assert "Deleted text mentioning" in body and "Anna" not in body  # tracked deletion scanned
        assert "mailto:priya" not in body and "Text box:" in body and "lars.eriksen" not in body  # field code, text box
        assert b"Footnote for" in parts["word/footnotes.xml"] and b"794" not in parts["word/footnotes.xml"]

    def test_metadata_and_hidden_parts_are_removed(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        _, _, out = docx_run
        assert out.data is not None
        names = set(fx.package_text(out.data))
        assert not {"docProps/thumbnail.jpeg", "docProps/custom.xml", "customXml/item1.xml"} & names
        view = fx.docx_text(out.data)
        assert view["meta"][0] == "" and view["meta"][1] == ""  # author, last modified by
        assert [c.author for c in view["doc"].comments] == ["Author"]
        assert b"thumbnail" not in b"".join(fx.package_text(out.data).values())

    def test_the_file_name_is_personal_data_too(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        assert docx_run[2].filename == "document"

    def test_redacting_twice_finds_nothing_more(self, docx_run) -> None:  # type: ignore[no-untyped-def]
        e, _, out = docx_run
        assert out.data is not None
        again = run(e, out.data, "x.docx")
        assert again.scan.counts == {} and again.verified

    def test_pictures_are_refused_by_default_and_blanked_on_request(self, env) -> None:  # type: ignore[no-untyped-def]
        data = fx.make_docx(picture=True)
        assert refused(env(), data, "a.docx") == "images_present"
        out = run(env(documents={"images": "blank"}), data, "a.docx")
        images = [n for n in fx.package_text(out.data or b"") if n.startswith("word/media/")]
        assert images and all(len(fx.package_text(out.data or b"")[n]) < 200 for n in images)  # a blank 1x1 image

    def test_text_only_mode_returns_text_without_a_file(self, env) -> None:  # type: ignore[no-untyped-def]
        out = run(env(), fx.make_docx(), "a.docx", want_file=False)
        assert out.data is None and "[PERSON_1]" in out.text and "priya.raman" not in out.text


# =========================================================================== Excel
class TestXlsx:
    def test_nothing_survives_in_any_part(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.verified and xlsx_run.data is not None
        assert fx.canaries_in_package(xlsx_run.data) == []

    def test_nothing_survives_for_an_independent_reader(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.data is not None
        joined = "\n".join(fx.xlsx_text(xlsx_run.data))
        for value in [*fx.CANARIES, "4111", "5500 0000"]:
            assert value not in joined, value

    def test_bare_names_are_found_through_their_column_header(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.data is not None
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(xlsx_run.data))["Customers"]
        assert ws["A2"].value.startswith("[PERSON_") and ws["F2"].value.startswith("[PERSON_")  # Name, Manager columns
        assert ws["A3"].value == ws["F2"].value  # the same person is the same placeholder everywhere

    def test_numbers_that_are_identifiers_become_text_placeholders(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.data is not None
        from openpyxl import load_workbook

        ws = load_workbook(io.BytesIO(xlsx_run.data))["Customers"]
        assert ws["E2"].value == "[NL_BSN_1]" and ws["E2"].data_type == "s"  # a BSN under a "BSN" header

    def test_hidden_sheet_names_formulas_and_names(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.data is not None
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(xlsx_run.data))
        assert wb["Hidden payroll"].sheet_state == "hidden" and wb["Hidden payroll"]["B2"].value.startswith("[US_SSN")
        assert "Priya Raman payroll" not in wb.sheetnames and any(n.startswith("Redacted") for n in wb.sheetnames)
        assert wb["Customers"]["H2"].value is None or "@" not in str(wb["Customers"]["H2"].value)
        assert "owner_mail" not in wb.defined_names

    def test_comments_and_authors(self, xlsx_run) -> None:  # type: ignore[no-untyped-def]
        assert xlsx_run.data is not None
        from openpyxl import load_workbook

        comment = load_workbook(io.BytesIO(xlsx_run.data))["Customers"]["A2"].comment
        assert comment is not None and "794" not in comment.text and comment.author == "Author"

    def test_too_many_cells_is_refused_not_truncated(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(documents={"max_segments": 10}), fx.make_xlsx(), "a.xlsx") == "too_many_segments"


# =========================================================================== PowerPoint
class TestPptx:
    def test_nothing_survives_in_any_part(self, pptx_run) -> None:  # type: ignore[no-untyped-def]
        assert pptx_run.verified and pptx_run.data is not None
        assert fx.canaries_in_package(pptx_run.data) == []

    def test_nothing_survives_for_an_independent_reader(self, pptx_run) -> None:  # type: ignore[no-untyped-def]
        assert pptx_run.data is not None
        joined = "\n".join(fx.pptx_text(pptx_run.data))
        for value in [*fx.CANARIES, "4111"]:
            assert value not in joined, value

    def test_notes_alt_text_and_split_runs(self, pptx_run) -> None:  # type: ignore[no-untyped-def]
        assert pptx_run.data is not None
        text = fx.pptx_text(pptx_run.data)
        assert any(t.startswith("Contact [PERSON_") for t in text)  # "Anna " + "Jansen" in two runs
        assert any(re.fullmatch(r"(Photo of )?\[PERSON_\d+\]", t) for t in text)  # alt text
        assert any("[CREDIT_CARD_1]" in t and "Ask [PERSON_" in t for t in text)  # speaker notes

    def test_reading_text_groups_by_slide(self, pptx_run) -> None:  # type: ignore[no-untyped-def]
        assert "[slide 1]" in pptx_run.text and "[slide 2]" in pptx_run.text and "[notes 1]" in pptx_run.text


# =========================================================================== PDF
class TestPdf:
    def test_every_sensitive_value_is_painted_black(self, pdf_run) -> None:  # type: ignore[no-untyped-def]
        original, out = pdf_run
        assert out.data is not None and out.verified
        fx.assert_blacked_out(original, out.data, [fx.EMAIL, fx.PHONE, fx.CARD, fx.IBAN, "Priya Raman", "Lars Eriksen"])

    def test_the_text_layer_is_gone_not_hidden(self, pdf_run) -> None:  # type: ignore[no-untyped-def]
        import pypdfium2 as pdfium

        original, out = pdf_run
        assert out.data is not None
        redacted, source = pdfium.PdfDocument(out.data), pdfium.PdfDocument(original)
        assert len(redacted) == len(source) == 3
        assert all(redacted[i].get_textpage().count_chars() == 0 for i in range(3))
        for needle in (b"priya", b"Priya", b"794", b"4111", b"NL91", b"Helvetica", b"/Info", b"/Producer"):
            assert needle not in out.data

    def test_page_sizes_and_unredacted_content_are_kept(self, pdf_run) -> None:  # type: ignore[no-untyped-def]
        import pypdfium2 as pdfium

        original, out = pdf_run
        assert out.data is not None
        a, b = pdfium.PdfDocument(original), pdfium.PdfDocument(out.data)
        for i in range(3):
            assert abs(a[i].get_size()[0] - b[i].get_size()[0]) < 1 and abs(a[i].get_size()[1] - b[i].get_size()[1]) < 1
        page3 = b[2].render(scale=1).to_pil().convert("L")
        assert page3.getextrema()[1] == 255  # the page is not simply black

    def test_pictures_and_pages_without_text_are_refused_by_default(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(), fx.make_pdf(picture=True), "a.pdf") == "images_present"
        assert refused(env(), fx.make_pdf(vector_only_page=True), "a.pdf") == "uninspectable_pages"

    def test_blank_mode_removes_pictures_and_blanks_uninspectable_pages(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env(documents={"images": "blank"})
        out = run(e, fx.make_pdf(picture=True, vector_only_page=True), "a.pdf")
        assert out.data is not None and out.info["blanked_pages"] == [3]
        import pypdfium2 as pdfium

        page = pdfium.PdfDocument(out.data)[2].render(scale=1).to_pil().convert("L")
        assert page.getextrema() == (255, 255)  # white

    def test_encrypted_truncated_and_empty_pdfs(self, env) -> None:  # type: ignore[no-untyped-def]
        from reportlab.lib.pdfencrypt import StandardEncryption
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, encrypt=StandardEncryption("user-pw", ownerPassword="owner-pw"))
        c.drawString(72, 700, "secret")
        c.save()
        e = env()
        assert refused(e, buf.getvalue(), "a.pdf") == "encrypted"
        assert refused(e, fx.make_pdf()[:300], "a.pdf") == "corrupt"
        assert refused(e, b"%PDF-1.4\n%%EOF\n", "a.pdf") == "corrupt"

    def test_a_large_page_is_rendered_at_a_lower_resolution_and_an_absurd_one_is_refused(self, env) -> None:  # type: ignore[no-untyped-def]
        from reportlab.pdfgen import canvas

        def page(size: int) -> bytes:
            buf = io.BytesIO()
            c = canvas.Canvas(buf, pagesize=(size, size))
            c.setFont("Helvetica", 40)
            c.drawString(100, 100, f"Contact {fx.EMAIL}")
            c.save()
            return buf.getvalue()

        out = run(
            env(), page(3000), "big.pdf"
        )  # 41 inches square: rendered below 150 dpi to stay within the pixel budget
        assert out.data is not None and len(out.data) < 5_000_000
        assert refused(env(), page(14000), "huge.pdf") == "page_too_large"

    def test_too_many_pages(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(documents={"max_pages": 2}), fx.make_pdf(), "a.pdf") == "too_many_pages"

    def test_a_key_is_redacted_and_a_policy_block_can_be_overridden(self, env) -> None:  # type: ignore[no-untyped-def]
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(72, 700, "deploy key: sk-proj-7Hq2LmV9xRt4KpZ8wNc3YbF6dJs1GaE5")
        c.drawString(72, 680, "refund to NL91ABNA0417164300")
        c.save()
        redacted = run(env(), buf.getvalue(), "a.pdf")  # keys are always redacted: the document is released
        assert redacted.data is not None and not redacted.scan.blocked and "SECRET" in redacted.scan.counts
        e = env(policy={"entities": {"IBAN": {"action": "block", "overridable": True}}})
        blocked = run(e, buf.getvalue(), "a.pdf")
        assert blocked.scan.blocked and blocked.data is None
        released = run(e, buf.getvalue(), "a.pdf", override=True)
        assert released.data is not None and not released.scan.blocked


# =========================================================================== text files
class TestText:
    def test_csv_names_are_found_through_headers_and_structure_is_kept(self, env) -> None:  # type: ignore[no-untyped-def]
        csv = f'Name,Email,Notes\r\nPriya Raman,{fx.EMAIL},"likes, commas"\r\nLars Eriksen,{fx.EMAIL2},ok\r\n'
        out = run(env(), csv.encode(), "people.csv")
        text = (out.data or b"").decode()
        assert text.split("\r\n")[0] == "Name,Email,Notes" and text.endswith("\r\n")
        assert 'PERSON_1],[EMAIL_1],"likes, commas"' in text and "Priya" not in text and "Eriksen" not in text
        assert out.content_type and out.extension == "csv"

    def test_json_stays_valid(self, env) -> None:  # type: ignore[no-untyped-def]
        doc = json.dumps({"user": {"email": fx.EMAIL, "card": fx.CARD, "note": "Call Priya Raman tomorrow"}, "n": 3})
        out = run(env(), doc.encode(), "data.json")
        parsed = json.loads(out.data or b"")
        assert parsed["n"] == 3 and parsed["user"]["email"] == "[EMAIL_1]" and "Priya" not in json.dumps(parsed)

    def test_encodings_and_byte_order_marks_are_preserved(self, env) -> None:  # type: ignore[no-untyped-def]
        for codec, bom in (("utf-8-sig", b"\xef\xbb\xbf"), ("utf-16", None)):
            data = f"Contact {fx.EMAIL} soon.".encode(codec)
            out = run(env(), data, "note.txt")
            assert out.data is not None and fx.EMAIL.encode(codec) not in out.data
            assert out.data.startswith(bom) if bom else out.data[:2] in (b"\xff\xfe", b"\xfe\xff")
            assert "[EMAIL_1]" in out.data.decode(codec)

    def test_keys_in_config_files_are_redacted_in_place(self, env) -> None:  # type: ignore[no-untyped-def]
        env_file = b"DEBUG=1\nOPENAI_API_KEY=sk-proj-7Hq2LmV9xRt4KpZ8wNc3YbF6dJs1GaE5\n"
        out = run(env(), env_file, ".env")
        assert not out.scan.blocked and out.scan.risk.value == "yellow"
        assert out.data == b"DEBUG=1\nOPENAI_API_KEY=[SECRET_1]\n"

    def test_unclear_files_are_refused(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env()
        assert refused(e, b"abc\x00def", "x.txt") == "binary_content"
        assert refused(e, "caf\xe9".encode("latin-1"), "x.txt") == "unsupported_encoding"
        assert refused(e, b"hello", "mystery.bin") == "unsupported_format"
        assert refused(e, b"hello", "x.docx") == "unsupported_format"  # not a zip, so not a docx whatever the name says

    def test_long_text_is_scanned_in_chunks_without_missing_the_boundary(self, env) -> None:  # type: ignore[no-untyped-def]
        body = ("filler words " * 700 + f"\nreach me at {fx.EMAIL}\n") * 3
        out = run(env(), body.encode(), "log.txt")
        assert out.data is not None and fx.EMAIL.encode() not in out.data and out.data.count(b"[EMAIL_1]") == 3


# =========================================================================== hostile files
class TestHostileFiles:
    def test_zip_bomb(self, env) -> None:  # type: ignore[no-untyped-def]
        bomb = with_part(fx.make_docx(), "word/big.xml", b"<a>" + b"A" * 30_000_000 + b"</a>")
        assert refused(env(), bomb, "a.docx") == "zip_bomb"

    def test_path_traversal_in_a_part_name(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(), with_part(fx.make_docx(), "../evil.xml", b"<a/>"), "a.docx") == "unsafe_path"
        assert refused(env(), with_part(fx.make_docx(), "word/..\\evil.xml", b"<a/>"), "a.docx") == "unsafe_path"

    def test_entity_declarations_are_refused(self, env) -> None:  # type: ignore[no-untyped-def]
        bomb = b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;">]><d>&b;</d>'
        assert refused(env(), with_part(fx.make_docx(), "word/document.xml", bomb), "a.docx") == "unsupported_xml"

    def test_macro_enabled_documents(self, env) -> None:  # type: ignore[no-untyped-def]
        data = fx.make_docx()
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            types = z.read("[Content_Types].xml").replace(b"document.main+xml", b"document.macroEnabled.main+xml")
        assert refused(env(), with_part(data, "[Content_Types].xml", types), "a.docm") == "macros"

    def test_embedded_objects_pivot_caches_and_external_links(self, env) -> None:  # type: ignore[no-untyped-def]
        assert (
            refused(env(), with_part(fx.make_docx(), "word/embeddings/oleObject1.bin", b"\x00data"), "a.docx")
            == "embedded_objects"
        )
        assert (
            refused(env(), with_part(fx.make_xlsx(), "xl/pivotCache/pivotCacheRecords1.xml", b"<a/>"), "a.xlsx")
            == "pivot_cache"
        )
        assert (
            refused(env(), with_part(fx.make_xlsx(), "xl/externalLinks/externalLink1.xml", b"<a/>"), "a.xlsx")
            == "external_links"
        )
        assert refused(env(), with_part(fx.make_docx(), "word/vbaProject.bin", b"\x00"), "a.docx") in (
            "macros",
            "unsupported_part",
        )

    def test_legacy_and_unrelated_files(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env()
        assert refused(e, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600, "old.doc") == "legacy_office"
        empty_zip = io.BytesIO()
        with zipfile.ZipFile(empty_zip, "w") as z:
            z.writestr("hello.txt", "hi")
        assert refused(e, empty_zip.getvalue(), "a.zip") == "not_office"
        assert refused(e, b"PK\x03\x04garbage", "a.docx") == "corrupt"

    def test_the_size_limit(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(documents={"max_bytes": 1_000}), b"x" * 2_000, "a.txt") == "too_large"

    def test_a_disabled_format(self, env) -> None:  # type: ignore[no-untyped-def]
        assert refused(env(documents={"formats": ["pdf"]}), fx.make_docx(), "a.docx") == "format_disabled"


# =========================================================================== the process boundary
class TestIsolation:
    def test_the_worker_runs_with_an_empty_environment_and_cannot_write_files(self, env) -> None:  # type: ignore[no-untyped-def]
        status = asyncio.run(env().runner.selftest())
        assert status["environment_empty"] and not status["can_write_files"]
        assert status["address_space_limited"] and status["cwd"] == "/" and status["cpu_limit_s"] == 60
        assert "network_isolated" in status

    def test_provider_keys_never_reach_the_worker(self, env, monkeypatch: pytest.MonkeyPatch) -> None:  # type: ignore[no-untyped-def]
        monkeypatch.setenv("OPENAI_API_KEY", "sk-should-never-be-visible")
        assert asyncio.run(env().runner.selftest())["environment_empty"]

    def test_a_stuck_worker_is_killed_and_the_next_one_works(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env()

        async def go() -> int | None:
            pid = None
            with pytest.raises(TimeoutError):
                async with asyncio.timeout(1), e.runner.session("text", "", b"") as session:
                    pid = session.proc.pid  # type: ignore[union-attr]
                    await session.call("sleep", seconds=30)
            assert pid is not None
            return pid

        pid = asyncio.run(go())
        with pytest.raises(ProcessLookupError):
            os.kill(pid, 0)  # the process is gone
        assert run(e, b"Contact priya.raman@corp.example", "a.txt").verified

    def test_a_crashing_worker_is_reported_not_hung(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env()

        async def go() -> str:
            async with e.runner.session("text", "a.txt", b"Contact priya.raman@corp.example") as session:
                os.killpg(session.proc.pid, signal.SIGKILL)  # type: ignore[union-attr]
                await asyncio.sleep(0.2)
                try:
                    await session.open({})
                except DocumentError as exc:
                    return exc.code
            return "no error"

        assert asyncio.run(go()) in ("worker_crashed", "worker_limit")

    def test_the_timeout_applies_to_the_whole_document(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env(documents={"timeout_s": 5})
        e.settings.documents.__dict__["timeout_s"] = 0.01  # below the configured minimum, only to test the guard
        assert refused(e, fx.make_docx(), "a.docx") == "timeout"

    def test_concurrent_documents_are_limited_and_all_succeed(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env(documents={"max_concurrent": 2})

        async def many() -> list[DocumentOutcome]:
            data = [
                fx.make_docx(),
                fx.make_xlsx(),
                fx.make_pptx(),
                fx.make_pdf(),
                b"Mail priya.raman@corp.example",
                fx.make_docx(),
            ]
            names = ["a.docx", "b.xlsx", "c.pptx", "d.pdf", "e.txt", "f.docx"]
            return await asyncio.gather(
                *(
                    redact_document(
                        e.pipeline,
                        e.settings.documents,
                        e.runner,
                        d,
                        n,
                        None,
                        mode=None,
                        context_repair=None,
                        override=False,
                        want_file=True,
                    )
                    for d, n in zip(data, names, strict=True)
                )
            )

        outcomes = asyncio.run(many())
        assert [o.kind for o in outcomes] == ["docx", "xlsx", "pptx", "pdf", "text", "docx"] and all(
            o.verified for o in outcomes
        )

    def test_the_queue_is_bounded_so_uploads_cannot_exhaust_memory(self, env) -> None:  # type: ignore[no-untyped-def]
        e = env(documents={"max_concurrent": 1})

        async def flood() -> list[object]:
            jobs = (
                redact_document(e.pipeline, e.settings.documents, e.runner, b"Mail priya.raman@corp.example", "a.txt", None,
                                mode=None, context_repair=None, override=False, want_file=True)
                for _ in range(24)
            )  # fmt: skip
            return await asyncio.gather(*jobs, return_exceptions=True)

        results = asyncio.run(flood())
        busy = [r for r in results if isinstance(r, DocumentError) and r.code == "busy"]
        done = [r for r in results if isinstance(r, DocumentOutcome)]
        assert busy and done and len(busy) + len(done) == 24  # the queue held 1 + 4; the rest were turned away
        assert e.runner.waiting == 0

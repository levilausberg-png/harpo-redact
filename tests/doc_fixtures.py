"""Documents seeded with personal data in the places it tends to hide."""

from __future__ import annotations

import io
import zipfile
from typing import Any

# Values that must not survive anywhere: not in any part of the file, not in any reader's view of it.
EMAIL = "priya.raman@corp.example"
EMAIL2 = "lars.eriksen@corp.example"
PHONE = "+31 20 794 6095"
CARD = "4111 1111 1111 1111"
IBAN = "NL91ABNA0417164300"
SSN = "123-45-6789"
BSN = "111222333"
NAMES = ["Priya Raman", "Lars Eriksen", "Anna Jansen"]
CANARIES = [
    EMAIL,
    EMAIL2,
    "priya.raman",
    "794 6095",
    "4111",
    "NL91ABNA",
    SSN,
    BSN,
    "Priya",
    "Raman",
    "Lars",
    "Eriksen",
    "Jansen",
]


def png_bytes() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (40, 20), (200, 30, 30)).save(out, "PNG")
    return out.getvalue()


def package_text(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def canaries_in_package(data: bytes, needles: list[str] | None = None) -> list[tuple[str, str]]:
    """Every (part, value) where a canary appears in the decompressed bytes of any part."""
    hits = []
    for name, blob in package_text(data).items():
        for needle in needles or CANARIES:
            if needle.encode() in blob or needle.encode("utf-16-le") in blob:
                hits.append((name, needle))
    return hits


def inject(
    data: bytes,
    parts: dict[str, bytes] | None = None,
    overrides: dict[str, str] | None = None,
    rels: dict[str, list[str]] | None = None,
) -> bytes:
    """Add parts, content-type overrides and relationships to an existing Office file."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            if info.filename in (parts or {}):
                continue  # replaced below
            blob = src.read(info.filename)
            if info.filename == "[Content_Types].xml" and overrides:
                extra = "".join(f'<Override PartName="/{p}" ContentType="{t}"/>' for p, t in overrides.items())
                blob = blob.replace(b"</Types>", extra.encode() + b"</Types>")
            if info.filename in (rels or {}):
                blob = blob.replace(b"</Relationships>", "".join(rels[info.filename]).encode() + b"</Relationships>")
            dst.writestr(info.filename, blob)
        for name, blob in (parts or {}).items():
            dst.writestr(name, blob)
    return out.getvalue()


# --------------------------------------------------------------------------- docx
def make_docx(*, picture: bool = False, hidden: bool = True) -> bytes:
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    cp = d.core_properties
    cp.author, cp.last_modified_by = "Priya Raman", "Lars Eriksen"
    cp.title, cp.subject, cp.keywords = "Quarterly notes", f"For {EMAIL2}", EMAIL
    cp.comments, cp.category = f"call {PHONE}", "HR"

    first = d.add_paragraph("Please contact ")
    first.add_run("Priya ")
    first.add_run("Raman").bold = True  # a name split across runs
    first.add_run(f" at {EMAIL} today.")
    d.add_paragraph(f"Phone: {PHONE}. Card {CARD}.")
    keep = d.add_paragraph("Formatting that must survive: ")
    keep.add_run("bold words").bold = True
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "IBAN", IBAN
    table.cell(1, 0).text, table.cell(1, 1).text = "Owner", "Lars Eriksen asked about this."
    section = d.sections[0]
    section.header.paragraphs[0].text = "Confidential: Lars Eriksen"
    section.footer.paragraphs[0].text = f"BSN {BSN}"
    d.add_comment(d.paragraphs[0].runs[0], text=f"Ask Lars Eriksen at {EMAIL2}", author="Lars Eriksen", initials="LE")
    if hidden:
        tracked = d.add_paragraph("Visible text. ")
        tracked._p.append(
            parse_xml(
                f'<w:del {nsdecls("w")} w:id="901" w:author="Priya Raman" w:date="2024-03-14T10:00:00Z">'
                '<w:r><w:delText xml:space="preserve">Deleted text mentioning Anna Jansen. </w:delText></w:r></w:del>'
            )
        )
        field = d.add_paragraph()
        field._p.append(
            parse_xml(
                f"<w:fldSimple {nsdecls('w')} w:instr='HYPERLINK \"mailto:{EMAIL}\"'>"
                "<w:r><w:t>write to us</w:t></w:r></w:fldSimple>"
            )
        )
        box = d.add_paragraph()
        box._p.append(
            parse_xml(
                f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape id="tb1" type="#_x0000_t202">'
                f"<v:textbox><w:txbxContent><w:p><w:r><w:t>Text box: {EMAIL2}</w:t></w:r></w:p></w:txbxContent></v:textbox>"
                "</v:shape></w:pict></w:r>"
            )
        )
    if picture:
        path = io.BytesIO(png_bytes())
        d.add_picture(path)
    buffer = io.BytesIO()
    d.save(buffer)
    footnotes = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<w:footnote w:id="1"><w:p><w:r><w:t>Footnote for Priya Raman, {PHONE}.</w:t></w:r></w:p></w:footnote></w:footnotes>'
    ).encode()
    custom = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
        b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Owner"><vt:lpwstr>Priya Raman</vt:lpwstr></property>'
        b"</Properties>"
    )
    return inject(
        buffer.getvalue(),
        parts={
            "word/footnotes.xml": footnotes,
            "docProps/custom.xml": custom,
            "docProps/thumbnail.jpeg": b"\xff\xd8\xff\xe0thumbnail of Priya Raman first page",
            "customXml/item1.xml": b"<root><owner>Priya Raman</owner></root>",
        },
        overrides={
            "word/footnotes.xml": "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            "docProps/custom.xml": "application/vnd.openxmlformats-officedocument.custom-properties+xml",
        },
        rels={
            "word/_rels/document.xml.rels": [
                '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>',
                f'<Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="mailto:{EMAIL}" TargetMode="External"/>',
            ],
            "_rels/.rels": [
                '<Relationship Id="rIdCust" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties" Target="docProps/custom.xml"/>',
                '<Relationship Id="rIdThumb" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/thumbnail" Target="docProps/thumbnail.jpeg"/>',
            ],
        },
    )


def docx_text(data: bytes) -> dict[str, Any]:
    """What an independent reader (python-docx) sees: nothing here comes from our engine."""
    import docx

    d = docx.Document(io.BytesIO(data))
    texts = [p.text for p in d.paragraphs]
    for t in d.tables:
        texts += [c.text for row in t.rows for c in row.cells]
    for s in d.sections:
        texts += [p.text for p in s.header.paragraphs + s.footer.paragraphs]
    texts += [c.text for c in d.comments] + [c.author for c in d.comments] + [c.initials or "" for c in d.comments]
    cp = d.core_properties
    meta = [cp.author, cp.last_modified_by, cp.title, cp.subject, cp.keywords, cp.comments, cp.category]
    return {"doc": d, "texts": texts, "meta": meta}


# --------------------------------------------------------------------------- xlsx
def make_xlsx() -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append(["Name", "Email", "Phone", "Card", "BSN", "Manager", "Notes"])
    ws.append(["Priya Raman", EMAIL, PHONE, CARD, int(BSN), "Lars Eriksen", "met at the conference"])
    ws.append(["Lars Eriksen", EMAIL2, "+31 20 794 6096", "5500 0000 0000 0004", 123456782, "Priya Raman", ""])
    ws["H2"] = f'=IF(A2="x","{EMAIL}","")'
    ws["A2"].comment = Comment(f"Call Priya Raman on {PHONE}", "Lars Eriksen")
    ws.add_table(Table(displayName="CustomerTable", ref="A1:G3"))
    hidden = wb.create_sheet("Hidden payroll")
    hidden.sheet_state = "hidden"
    hidden.append(["Employee", "SSN"])
    hidden.append(["Anna Jansen", SSN])
    named = wb.create_sheet("Priya Raman payroll")
    named.append(["Note"])
    named.append(["nothing secret here"])
    wb.defined_names["owner_mail"] = DefinedName("owner_mail", attr_text=f'"{EMAIL}"')
    wb.properties.creator, wb.properties.lastModifiedBy = "Priya Raman", "Lars Eriksen"
    wb.properties.title, wb.properties.description = "Customer list", f"owner {EMAIL}"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def xlsx_text(data: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    out = list(wb.sheetnames) + [str(v) for v in wb.defined_names.values()]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    out.append(str(cell.value))
                if cell.comment:
                    out += [cell.comment.text, cell.comment.author]
        out += [t.displayName for t in ws.tables.values()]
    p = wb.properties
    out += [p.creator or "", p.lastModifiedBy or "", p.title or "", p.description or ""]
    return out


# --------------------------------------------------------------------------- pptx
def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Meeting with Priya Raman"
    slide.placeholders[1].text = f"Email {EMAIL} or call {PHONE}"
    slide.notes_slide.notes_text_frame.text = f"Card {CARD}. Ask Lars Eriksen about it."
    box = slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    para = box.text_frame.paragraphs[0]
    para.add_run().text = "Contact Anna "
    para.add_run().text = "Jansen"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(6), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "IBAN", IBAN
    shape = slide.shapes.add_shape(1, Inches(6), Inches(5), Inches(1), Inches(1))
    shape._element.nvSpPr.cNvPr.set("descr", "Photo of Priya Raman")
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = f"BSN {BSN} for Lars Eriksen"
    prs.core_properties.author = "Priya Raman"
    prs.core_properties.title = f"Deck for {EMAIL2}"
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def pptx_text(data: bytes) -> list[str]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    out = [prs.core_properties.author, prs.core_properties.title]
    for slide in prs.slides:
        for shape in slide.shapes:
            out.append(shape._element.xpath("string(.//p:cNvPr/@descr)"))
            if shape.has_text_frame:
                out += [p.text for p in shape.text_frame.paragraphs]
            if shape.has_table:
                out += [c.text for row in shape.table.rows for c in row.cells]
        if slide.has_notes_slide:
            out.append(slide.notes_slide.notes_text_frame.text)
    return out


# --------------------------------------------------------------------------- pdf
def make_pdf(*, picture: bool = False, vector_only_page: bool = False) -> bytes:
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas

    out = io.BytesIO()
    c = canvas.Canvas(out, pagesize=(612, 792))
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 740, "Invoice for Priya Raman")
    c.setFont("Helvetica", 12)
    c.drawString(72, 710, f"Email: {EMAIL}")
    c.drawString(72, 690, f"Phone: {PHONE}")
    c.drawString(72, 670, f"Card: {CARD}")
    c.drawString(72, 650, "Nothing sensitive on this line, just ordinary words.")
    if picture:
        c.drawImage(ImageReader(io.BytesIO(png_bytes())), 72, 500, 100, 50)
    c.showPage()
    c.setPageRotation(90)  # a rotated page
    c.setFont("Helvetica", 12)
    c.drawString(72, 700, f"IBAN {IBAN} for Lars Eriksen")
    c.showPage()
    if vector_only_page:
        c.rect(100, 100, 200, 200, fill=1)
    else:
        c.setFont("Helvetica", 12)
        c.drawString(72, 700, "A page with no personal data at all.")
    c.showPage()
    c.save()
    return out.getvalue()


def assert_blacked_out(original: bytes, redacted: bytes, needles: list[str]) -> None:
    """Every needle's original position is solid black in the redacted PDF, page by page."""
    import pypdfium2 as pdfium

    src, out = pdfium.PdfDocument(original), pdfium.PdfDocument(redacted)
    for needle in needles:
        located = False
        for i in range(len(src)):
            textpage = src[i].get_textpage()
            text = textpage.get_text_range()
            k = text.find(needle)
            if k < 0:
                continue
            located = True
            boxes = [textpage.get_charbox(j) for j in range(k, k + len(needle))]
            left, bottom = min(b[0] for b in boxes), min(b[1] for b in boxes)
            right, top = max(b[2] for b in boxes), max(b[3] for b in boxes)
            # Boxes are in the SOURCE page's coordinates (before rotation): convert with the source page, then
            # look at the same pixels of the redacted page, which is already in its final orientation.
            source_page = src[i]
            convert = source_page.render(scale=2).get_posconv(source_page)
            (x0, y0), (x1, y1) = convert.to_bitmap(left, top), convert.to_bitmap(right, bottom)
            bitmap = out[i].render(scale=2)
            region = (
                bitmap.to_pil().convert("L").crop((min(x0, x1) + 2, min(y0, y1) + 2, max(x0, x1) - 2, max(y0, y1) - 2))
            )
            pixels = list(region.tobytes())  # 8-bit greyscale
            assert pixels and max(pixels) < 60, (
                f"{needle!r} on page {i + 1} is not fully black (brightest pixel {max(pixels or [0])})"
            )
        assert located, f"{needle!r} not found in the original"
